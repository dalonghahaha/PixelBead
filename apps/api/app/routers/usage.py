"""spec 011 — 用量清单端点(JSON / Excel / CSV)

GET /patterns/{id}/usage        → JSON
GET /patterns/{id}/usage.xlsx   → Excel(.xlsx, 流式返回)
GET /patterns/{id}/usage.csv    → CSV(UTF-8 BOM)
"""
from __future__ import annotations

import io
import logging
from collections import Counter
from typing import Any

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from ..db import get_db
from ..deps import get_current_user
from ..models import Pattern, User
from ..services import pypindou_service

router = APIRouter(prefix="/patterns", tags=["patterns"])
log = logging.getLogger(__name__)


# ---- Schemas ----


class UsageItem(BaseModel):
    code: str
    rgb: str
    name_zh: str | None = None
    name_en: str | None = None
    count: int
    packs: int


class UsageReport(BaseModel):
    pattern_id: str
    bead_size: str = "mini"
    items: list[UsageItem]
    total_count: int
    total_packs: int
    generated_at: str


# ---- 计算逻辑(共享) ----


def _calc_usage(pattern: Pattern, db: Session, beads_per_pack: int = 500) -> dict[str, Any]:
    """从 Pattern.color_counts 计算用量清单

    color_counts 形如:{"H1": 256, "H2": 89, ...}
    返回:UsageReport dict
    """
    if not pattern.color_counts:
        raise ValueError("图纸 color_counts 为空,无法生成用量清单")

    # 拉取色板完整信息(用于 name_zh / name_en)
    palettes = pypindou_service.list_palettes_safe()
    palette_meta = next((p for p in palettes if p["id"] == pattern.palette), {})
    color_lookup = {c["code"]: c for c in palette_meta.get("colors", [])}

    # 计算每色号用量 + 包数(向上取整,count=0 → packs=0)
    items: list[dict[str, Any]] = []
    for code, count in pattern.color_counts.items():
        meta = color_lookup.get(code, {})
        rgb_tuple = meta.get("rgb", [255, 255, 255])
        rgb_hex = "#{:02X}{:02X}{:02X}".format(*rgb_tuple)
        items.append(
            {
                "code": code,
                "rgb": rgb_hex,
                "name_zh": meta.get("name"),
                "name_en": meta.get("name_en"),
                "count": count,
                "packs": 0 if count == 0 else -(-count // beads_per_pack),  # ceil
            }
        )

    # 按颗数降序
    items.sort(key=lambda x: x["count"], reverse=True)

    from datetime import datetime

    return {
        "pattern_id": pattern.id,
        "bead_size": pattern.bead_size or "mini",
        "items": items,
        "total_count": sum(i["count"] for i in items),
        "total_packs": sum(i["packs"] for i in items),
        "generated_at": datetime.utcnow().isoformat() + "Z",
    }


# ---- 端点 ----


@router.get("/{pattern_id}/usage", response_model=UsageReport)
async def get_usage(
    pattern_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
    beads_per_pack: int = 500,
):
    """JSON 用量清单"""
    pattern = db.get(Pattern, pattern_id)
    if not pattern:
        raise HTTPException(status_code=404, detail="图纸不存在")
    if pattern.user_id != current_user.id:
        raise HTTPException(status_code=403, detail="无权访问")

    try:
        return _calc_usage(pattern, db, beads_per_pack)
    except ValueError as e:
        raise HTTPException(status_code=409, detail=str(e))


@router.get("/{pattern_id}/usage.xlsx")
async def get_usage_xlsx(
    pattern_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
    beads_per_pack: int = 500,
):
    """Excel(.xlsx)用量清单"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    pattern = db.get(Pattern, pattern_id)
    if not pattern:
        raise HTTPException(status_code=404, detail="图纸不存在")
    if pattern.user_id != current_user.id:
        raise HTTPException(status_code=403, detail="无权访问")

    try:
        report = _calc_usage(pattern, db, beads_per_pack)
    except ValueError as e:
        raise HTTPException(status_code=409, detail=str(e))

    wb = Workbook()
    ws = wb.active
    ws.title = "用量清单"

    # 文件头
    ws["A1"] = f"PixelBead 用量清单 - {pattern.id[:8]} - {pattern.bead_size}"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:F1")

    # 列头
    headers = ["色号", "色块", "中文名", "英文名", "颗数", "包数"]
    ws.append([])  # 空行
    ws.append(headers)
    for cell in ws[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 数据行
    for item in report["items"]:
        row_idx = ws.max_row + 1
        ws.cell(row=row_idx, column=1, value=item["code"])
        # 色块单元格(列 B,空字符串 + 背景色)
        color_cell = ws.cell(row=row_idx, column=2, value="")
        rgb_hex = item["rgb"].lstrip("#")
        color_cell.fill = PatternFill("solid", fgColor=rgb_hex)
        ws.cell(row=row_idx, column=3, value=item["name_zh"] or "")
        ws.cell(row=row_idx, column=4, value=item["name_en"] or "")
        ws.cell(row=row_idx, column=5, value=item["count"]).alignment = Alignment(horizontal="right")
        ws.cell(row=row_idx, column=6, value=item["packs"]).alignment = Alignment(horizontal="right")

    # 汇总行
    ws.append([])
    ws.cell(row=ws.max_row + 1, column=1, value="合计").font = Font(bold=True)
    ws.cell(row=ws.max_row, column=5, value=report["total_count"]).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=6, value=report["total_packs"]).font = Font(bold=True)

    # 列宽
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 6
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 10
    ws.column_dimensions["F"].width = 10

    # 流式返回
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    from datetime import datetime

    filename = f"pixelbead-usage-{pattern.id[:8]}-{datetime.utcnow().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/{pattern_id}/usage.csv")
async def get_usage_csv(
    pattern_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
    beads_per_pack: int = 500,
):
    """CSV 用量清单(UTF-8 BOM,Excel 中文兼容)"""
    pattern = db.get(Pattern, pattern_id)
    if not pattern:
        raise HTTPException(status_code=404, detail="图纸不存在")
    if pattern.user_id != current_user.id:
        raise HTTPException(status_code=403, detail="无权访问")

    try:
        report = _calc_usage(pattern, db, beads_per_pack)
    except ValueError as e:
        raise HTTPException(status_code=409, detail=str(e))

    # UTF-8 BOM
    buf = io.StringIO()
    buf.write("\ufeff")
    # 列头
    buf.write("色号,中文名,英文名,颗数,包数\n")
    for item in report["items"]:
        # CSV 转义:逗号 / 引号
        def esc(s: str | None) -> str:
            if s is None:
                return ""
            if any(c in s for c in (",", '"', "\n")):
                return '"' + s.replace('"', '""') + '"'
            return s

        buf.write(
            f"{esc(item['code'])},{esc(item['name_zh'])},{esc(item['name_en'])},{item['count']},{item['packs']}\n"
        )

    from datetime import datetime

    filename = f"pixelbead-usage-{pattern.id[:8]}-{datetime.utcnow().strftime('%Y%m%d')}.csv"
    return StreamingResponse(
        iter([buf.getvalue().encode("utf-8")]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )